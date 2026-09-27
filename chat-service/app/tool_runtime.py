"""Bounded, role-aware tools for the chat model."""
from __future__ import annotations

import asyncio
import base64
import csv
import io
import ipaddress
import json
import re
import socket
import ssl
from html.parser import HTMLParser
from http.client import HTTPSConnection
from urllib.parse import parse_qs, quote_plus, urljoin, urlsplit
from typing import Any, Literal

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

WEB_ROLES = {"ventas", "admin"}
MAX_WEB_BYTES = 100_000
MAX_DOCUMENT_CHARS = 8_000
INJECTION = re.compile(
    r"(?i)(ignore|ignora|olvida|desobedece).{0,70}(instruction|instrucci[oó]n|regla|system|sistema)"
    r"|(?:system|developer|assistant)\s*:\s*|<\s*/?\s*(?:system|developer|assistant)\s*>"
    r"|(?:api[_ -]?key|password|contrase[ñn]a|token).{0,50}(?:send|env[ií]a|exfiltra)"
)


def capabilities(roles: list[str], authenticated: bool, business_tools: list,
                 documents: list[dict] | None = None) -> list[dict]:
    if not authenticated or "guest" in roles:
        return []
    specs = [
        {"type": "function", "function": {"name": item.name,
            "description": item.description,
            "parameters": item.input_schema}}
        for item in business_tools
    ]
    specs.append({"type": "function", "function": {"name": "ask_user",
        "description": "Muestra un cuadro de diálogo para pedir un dato necesario al usuario. No inventes su respuesta.",
        "parameters": {"type": "object", "properties": {
            "question": {"type": "string"}, "field_label": {"type": "string"}},
            "required": ["question", "field_label"], "additionalProperties": False}}})
    document_ids = [doc["id"] for doc in (documents or []) if isinstance(doc, dict) and isinstance(doc.get("id"), str)][:30]
    if document_ids:
        specs.append({"type": "function", "function": {"name": "attach_existing_document",
            "description": "Adjunta en el chat un documento existente autorizado para este usuario.",
            "parameters": {"type": "object", "properties": {"document_id": {
                "type": "string", "enum": document_ids}},
                "required": ["document_id"], "additionalProperties": False}}})
    if WEB_ROLES.intersection(roles):
        specs.extend([
            {"type": "function", "function": {"name": "web_search",
                "description": "Busca información pública actual en internet. No incluyas datos internos, personales ni secretos en la consulta.",
                "parameters": {"type": "object", "properties": {"query": {"type": "string"}},
                    "required": ["query"], "additionalProperties": False}}},
            {"type": "function", "function": {"name": "web_open",
                "description": "Lee una página HTTPS pública. Nunca envíes datos internos en la URL.",
                "parameters": {"type": "object", "properties": {"url": {"type": "string"}},
                    "required": ["url"], "additionalProperties": False}}},
        ])
    return specs


def safe_public_url(url: str) -> tuple[str, str, str]:
    if not isinstance(url, str) or len(url) > 2048:
        raise ValueError("URL inválida")
    parsed = urlsplit(url)
    if parsed.scheme != "https" or not parsed.hostname or parsed.port not in (None, 443):
        raise ValueError("Solo se permiten páginas HTTPS públicas")
    if parsed.username or parsed.password or parsed.fragment:
        raise ValueError("URL no permitida")
    host = parsed.hostname.rstrip(".").lower()
    if host == "localhost" or host.endswith(".localhost") or host.endswith(".local"):
        raise ValueError("Host no público")
    addresses = {entry[4][0] for entry in socket.getaddrinfo(host, 443, type=socket.SOCK_STREAM)}
    if not addresses or any(not ipaddress.ip_address(ip).is_global for ip in addresses):
        raise ValueError("Dirección no pública")
    return host, sorted(addresses)[0], (parsed.path or "/") + ("?" + parsed.query if parsed.query else "")


class _PinnedHTTPS(HTTPSConnection):
    def __init__(self, host: str, address: str):
        super().__init__(host, 443, timeout=6)
        self.address = address

    def connect(self):
        raw = socket.create_connection((self.address, 443), timeout=self.timeout)
        self.sock = ssl.create_default_context().wrap_socket(raw, server_hostname=self.host)


def _fetch(url: str) -> tuple[str, str]:
    for _ in range(3):
        host, address, path = safe_public_url(url)
        connection = _PinnedHTTPS(host, address)
        try:
            connection.request("GET", path, headers={"Host": host, "User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 Chrome/120.0 Safari/537.36", "Accept": "text/html,text/plain"})
            response = connection.getresponse()
            if response.status in (301, 302, 303, 307, 308):
                location = response.getheader("Location")
                if not location:
                    raise ValueError("Redirección inválida")
                url = urljoin(url, location)
                continue
            if response.status != 200:
                raise ValueError(f"La página respondió HTTP {response.status}")
            content_type = response.getheader("Content-Type", "").lower()
            if not ("text/html" in content_type or "text/plain" in content_type):
                raise ValueError("Tipo de contenido no admitido")
            raw = response.read(MAX_WEB_BYTES + 1)
            return url, raw[:MAX_WEB_BYTES].decode("utf-8", errors="replace")
        finally:
            connection.close()
    raise ValueError("Demasiadas redirecciones")


class _PageText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.hidden = 0
        self.parts = []
        self.links = []
        self.current = None

    def handle_starttag(self, tag, attrs):
        data = dict(attrs)
        if tag in ("script", "style", "noscript"):
            self.hidden += 1
        if tag == "a" and "result__a" in data.get("class", ""):
            self.current = {"url": data.get("href", ""), "title": ""}

    def handle_endtag(self, tag):
        if tag in ("script", "style", "noscript") and self.hidden:
            self.hidden -= 1
        if tag == "a" and self.current:
            self.links.append(self.current)
            self.current = None

    def handle_data(self, data):
        if not self.hidden:
            value = data.strip()
            if value:
                self.parts.append(value)
                if self.current:
                    self.current["title"] += value


async def web_search(query: str) -> dict:
    if not isinstance(query, str) or not 2 <= len(query.strip()) <= 180 or INJECTION.search(query):
        raise ValueError("Consulta no permitida")
    url, html = await asyncio.to_thread(_fetch, "https://html.duckduckgo.com/html/?q=" + quote_plus(query))
    parser = _PageText()
    parser.feed(html)
    results = []
    for link in parser.links[:8]:
        href = link["url"]
        if href.startswith("//"):
            href = "https:" + href
        if "uddg=" in href:
            href = parse_qs(urlsplit(href).query).get("uddg", [href])[0]
        if urlsplit(href).scheme == "https" and not INJECTION.search(link["title"]):
            results.append({"title": link["title"][:180], "url": href[:2048]})
    return {"query": query, "results": results[:5], "source": url}


async def web_open(url: str) -> dict:
    final_url, html = await asyncio.to_thread(_fetch, url)
    parser = _PageText()
    parser.feed(html)
    text = " ".join(parser.parts)[:6000]
    if INJECTION.search(text):
        raise ValueError("La página contiene instrucciones sospechosas")
    return {"url": final_url, "content": text, "untrusted_source": True}


def make_document(title: str, content: str, format: Literal["txt", "csv", "pdf", "docx", "xlsx"]) -> dict[str, Any]:
    if not isinstance(title, str) or not isinstance(content, str) or format not in ("txt", "csv", "pdf", "docx", "xlsx"):
        raise ValueError("Documento inválido")
    if not 1 <= len(title) <= 100 or not 1 <= len(content) <= MAX_DOCUMENT_CHARS:
        raise ValueError("Documento fuera de tamaño permitido")
    clean_title = re.sub(r"[^\w\- ]", "", title, flags=re.UNICODE).strip()[:60] or "documento"
    if format == "pdf":
        output = io.BytesIO()
        page = canvas.Canvas(output, pagesize=A4)
        page.setFont("Helvetica-Bold", 14)
        page.drawString(40, 800, clean_title)
        page.setFont("Helvetica", 10)
        y = 775
        for line in content.splitlines():
            for part in [line[i:i+95] for i in range(0, len(line), 95)] or [""]:
                if y < 45:
                    page.showPage()
                    page.setFont("Helvetica", 10)
                    y = 800
                page.drawString(40, y, part)
                y -= 15
        page.save()
        data = output.getvalue()
        mime = "application/pdf"
    elif format == "docx":
        output = io.BytesIO()
        document = Document()
        document.add_heading(clean_title, level=0)
        for paragraph in content.splitlines():
            document.add_paragraph(paragraph)
        document.save(output)
        data = output.getvalue()
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif format == "xlsx":
        output = io.BytesIO()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Datos"
        rows = list(csv.reader(io.StringIO(content)))
        if len(rows) > 200 or any(len(row) > 25 for row in rows):
            raise ValueError("Hoja de cálculo fuera de tamaño permitido")
        for row in rows:
            safe = ["'" + cell if cell.lstrip().startswith(("=", "+", "-", "@"))
                    else cell for cell in row]
            sheet.append(safe)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="334155")
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(42, max(12, max(
                len(str(cell.value or "")) for cell in column) + 2))
        workbook.save(output)
        data = output.getvalue()
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif format == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        for line in content.splitlines():
            writer.writerow([cell if not cell.lstrip().startswith(("=", "+", "-", "@")) else "'" + cell for cell in line.split(",")])
        data = output.getvalue().encode("utf-8-sig")
        mime = "text/csv"
    else:
        data = content.encode("utf-8")
        mime = "text/plain"
    return {"type": "attachment", "name": f"{clean_title}.{format}", "mime": mime,
            "data": base64.b64encode(data).decode("ascii")}
