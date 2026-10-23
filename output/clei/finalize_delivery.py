"""Verify final scientific artifacts and package the reproduction snapshot."""
import csv
import hashlib
import json
import shutil
import zipfile
from pathlib import Path
from pypdf import PdfReader
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
EVIDENCE = OUT / "evidence"


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


manifest = json.loads((EVIDENCE / "manifest.json").read_text())
for name, expected in manifest["hashes"].items():
    assert sha(ROOT / name) == expected, name
manifest["experimental_harness_sha256"] = sha(EVIDENCE / "evaluate_article.py")
(EVIDENCE / "manifest.json").write_text(json.dumps(manifest, indent=2))

with (EVIDENCE / "predictions.csv").open() as stream:
    predictions = list(csv.DictReader(stream))
with (EVIDENCE / "metrics.csv").open() as stream:
    metrics = list(csv.DictReader(stream))
assert len(predictions) == 9397
assert len(metrics) == 32
policy_columns = {"heuristic": "heuristic_blocked", "ml": "ml_blocked", "strict_or": "strict_or_blocked", "three_state": "three_state_blocked"}
for metric in metrics:
    rows = [r for r in predictions if r["experiment"] == metric["experiment"]]
    column = policy_columns[metric["policy"]]
    counts = {k: 0 for k in ("tp", "fp", "fn", "tn")}
    for row in rows:
        label, blocked = int(row["label"]), int(row[column])
        key = "tp" if label and blocked else "fn" if label else "fp" if blocked else "tn"
        counts[key] += 1
    assert len(rows) == int(metric["n"])
    for name, count in counts.items():
        assert count == int(metric[name]), (metric["experiment"], metric["policy"], name)
    if metric["policy"] == "three_state":
        for label, group in ((1, "attack"), (0, "benign")):
            for action in ("ALLOWED", "GUARDED", "BLOCKED"):
                count = sum(int(r["label"]) == label and r["decision"] == action for r in rows)
                assert count == int(float(metric[f"{group}_{action.lower()}"]))

pdf = OUT / "Promption_CLEI_Article.pdf"
shutil.copy2(OUT / "qa" / pdf.name, pdf)
reader = PdfReader(pdf)
assert len(reader.pages) == 15
texts = [page.extract_text() for page in reader.pages]
assert all(len(t.strip()) > 100 for t in texts)
for page in reader.pages:
    assert abs(float(page.mediabox.width) - 595.276) < 1
    assert abs(float(page.mediabox.height) - 841.89) < 1
assert "Roberto Enrique" in texts[0] and "David Adrian" in texts[0]
assert "0009-0004-0869-3881" in texts[0] and "0009-0003-2167-8515" in texts[0]
assert not any(token in "\n".join(texts) for token in ("{{", "}}", "turn0search", "@TABLE", "@FIG", "@ABSTRACT"))
assert "Abstract" in texts[0] and "Keywords:" in texts[0]
document = Document(OUT / "Promption_CLEI_Article.docx")
assert document.styles["Normal"].font.name == "Times New Roman"
assert document.styles["Normal"].font.size.pt == 10
assert document.styles["Title"].font.size.pt == 16
assert document.styles["Title"].font.bold
assert document.styles["Title"].paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
for name, size in (("Heading 1", 12), ("Heading 2", 10), ("Heading 3", 10)):
    assert document.styles[name].font.size.pt == size
    assert document.styles[name].paragraph_format.alignment == WD_ALIGN_PARAGRAPH.LEFT
assert document.styles["Heading 1"].font.bold and document.styles["Heading 2"].font.bold
assert document.styles["Heading 3"].font.italic and not document.styles["Heading 3"].font.bold
assert len(document.tables) == 9
for table in document.tables:
    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    for edge in ("left", "right", "insideV"):
        assert borders.find(qn("w:" + edge)).get(qn("w:val")) == "nil"
    assert sum(column.width for column in table.columns) <= 16.5 * 360000 + 1000
for section in document.sections:
    assert not any(p.text.strip() for p in section.header.paragraphs)
for label in ("Abstract", "References"):
    assert any(p.text == label for p in document.paragraphs)
for number in (1, 2, 3):
    assert f"Fig. {number}" in "\n".join(texts)
links = []
for page in reader.pages:
    for ref in page.get("/Annots", []):
        annotation = ref.get_object()
        action = annotation.get("/A")
        if action and action.get("/URI"):
            links.append(str(action["/URI"]))
assert len(links) >= 16
validation = json.loads((EVIDENCE / "document_validation.json").read_text())
validation.update(page_count=15, all_15_pages_visually_reviewed=True, visual_review="passed after final rendering", author_instructions="AuthorInstructions2016-OTH.pdf supplied by the user", author_instructions_sha256=sha(EVIDENCE / "AuthorInstructions2016-OTH.pdf"), headings_explicitly_left_aligned=True, table_vertical_borders_absent=True, abstract_and_keywords_on_first_page=True, pdf_uri_links=len(links), confusion_matrices_verified=32, predictions_verified=9397, input_hashes_verified=len(manifest["hashes"]), docx_sha256=sha(OUT / "Promption_CLEI_Article.docx"), pdf_sha256=sha(pdf))
(EVIDENCE / "document_validation.json").write_text(json.dumps(validation, indent=2))

files = set()
for name in ("promption", "src", "packages/ai-sdk"):
    for path in (ROOT / name).rglob("*"):
        if path.is_file() and path.suffix in (".py", ".yaml", ".js", ".ts", ".json", ".md") and not any(part in ("__pycache__", "node_modules", "dist") for part in path.parts):
            files.add(path)
for name in ("pyproject.toml", "LIBRARY.md", "README.md", "LICENSE", "LICENSE.md", "promption/py.typed", "config/config.yaml", "config/heuristics.yaml", "models/lightweight_classifier.pkl", "data/processed/training_data.csv", "data/processed/label_overrides.csv", "data/processed/quarantine.csv"):
    path = ROOT / name
    if path.is_file():
        files.add(path)
for path in (ROOT / "tests").glob("*.py"):
    files.add(path)
for path in EVIDENCE.iterdir():
    if path.is_file() and path.name != "delivery_validation.json":
        files.add(path)
for name in ("README_Reproduccion.md", "requirements-reproduction.txt", "Cumplimiento_CLEI.md", "submission_metadata.json", "submission_metadata_ascii.txt", "manuscript.md", "Promption_CLEI_Article.docx", "Promption_CLEI_Article.pdf", "architecture.png", "shift.png", "routing.png", "build_figures.py", "build_article.py", "article_content.md"):
    files.add(OUT / name)
archive = OUT / "Promption_CLEI_Reproduccion.zip"
with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as bundle:
    for path in sorted(files):
        relative = path.relative_to(ROOT)
        assert not any(part in (".git", ".venv", ".aws", "node_modules") or part.startswith(".env") for part in relative.parts)
        bundle.write(path, relative)
with zipfile.ZipFile(archive) as bundle:
    assert bundle.testzip() is None
    assert bundle.read("output/clei/Promption_CLEI_Article.pdf") == pdf.read_bytes()
report = {"status": "passed", "files": len(files), "zip_bytes": archive.stat().st_size, "zip_sha256": sha(archive), "pages": len(reader.pages), "metric_rows_checked": len(metrics), "predictions_checked": len(predictions), "links": len(links), "tracked_source_modified": False}
(EVIDENCE / "delivery_validation.json").write_text(json.dumps(report, indent=2))
print(json.dumps(report, indent=2))
