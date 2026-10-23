"""Validate the deployment locally without calling an LLM provider."""
import asyncio
import base64
import http.client
import io
import json
import subprocess
import sys
import time
from pathlib import Path


def request(port, path, method="GET", payload=None, cookie="", headers=None):
    connection = http.client.HTTPConnection("127.0.0.1", port, timeout=30)
    outgoing = {"Content-Type": "application/json", **(headers or {})}
    if cookie:
        outgoing["Cookie"] = cookie
    connection.request(method, path, json.dumps(payload) if payload is not None else None, outgoing)
    response = connection.getresponse()
    result = response.status, response.getheaders(), response.read()
    connection.close()
    return result


def verify_http():
    for port, path in ((8100, "/api/v1/health"), (8101, "/health"), (3100, "/")):
        status, _, _ = request(port, path)
        assert status == 200, (port, status)
    assert request(8100, "/api/v1/filter", "POST", {"text": "hola"})[0] == 401
    unauthenticated_request = {
        "text": "dime qué puedes hacer",
        "user": {"id": "CLI-CUST-01", "name": "Prueba", "email": "cliente@demo.shop",
                 "roles": ["customer"], "authenticated": True},
        "context": {},
    }
    assert request(8101, "/api/v1/chat", "POST", unauthenticated_request)[0] in (401, 403)
    assert request(3100, "/api/ai/turn", "POST", {})[0] in (401, 403)
    assert request(3100, "/api/admin/log-stats")[0] == 403
    print("Salud y endpoints privados: correctos")

    for email, role in (("cliente@demo.shop", "customer"), ("ana@demo.shop", "ventas"), ("jefe@demo.shop", "admin")):
        status, headers, body = request(3100, "/api/login", "POST", {"email": email, "password": "demo123"})
        assert status == 200
        cookies = [v for k, v in headers if k.lower() == "set-cookie"]
        assert cookies and all("Secure" in v and "HttpOnly" in v for v in cookies)
        cookie = "; ".join(v.split(";", 1)[0] for v in cookies)
        status, _, body = request(3100, "/api/chat", "POST", {"text": "dime qué puedes hacer"}, cookie)
        assert status == 200
        data = json.loads(body)
        assert not data["blocked"], (role, data.get("reason"))
        assert data.get("model") == "Promption", (role, data.get("model"))
        assert data.get("execution_metrics", {}).get("provider_calls", 0) == 0
        assert bool("Buscar información pública en internet" in data["reply"]) == (role != "customer")
        status, _, body = request(3100, "/api/chat", cookie=cookie)
        assert status == 200 and len(json.loads(body)["messages"]) >= 2
        status, _, body = request(3100, "/api/chat", "POST", {"text": "dime qué puedes hacer"}, cookie,
                                  {"Accept": "text/event-stream"})
        assert status == 200 and b"data:" in body
        events = [json.loads(line[6:]) for line in body.decode().splitlines() if line.startswith("data: ")]
        assert events and events[-1]["type"] == "result"
        assert not events[-1]["data"]["blocked"]
        print(f"Sesión {role}: capacidades, permisos, historial y streaming correctos")

    if "--restart-check" in sys.argv:
        before = json.loads(request(3100, "/api/chat", cookie=cookie)[2])["messages"]
        subprocess.run(["systemctl", "--user", "restart", "promption-chat"], check=True)
        for _ in range(50):
            try:
                if request(8101, "/health")[0] == 200:
                    break
            except OSError:
                pass
            time.sleep(.2)
        else:
            raise RuntimeError("El chat no respondió después de reiniciarlo")
        status, _, body = request(3100, "/api/chat", cookie=cookie)
        assert status == 200 and json.loads(body)["messages"] == before
        print("SQLite: historial conservado después de reiniciar el chat")


async def verify_mcp():
    from app.mcp_tools import MCPToolExecutor
    from docx import Document
    from openpyxl import load_workbook

    executor = MCPToolExecutor()
    for format in ("docx", "pdf", "xlsx"):
        result = await executor.execute("make_document", {
            "title": "Prueba del despliegue", "content": "Nombre,Valor\nDato,=1+1", "format": format,
        }, ["customer"])
        assert result["audit"]["allowed"]
        raw = base64.b64decode(result["result"]["data"])
        if format == "docx":
            assert Document(io.BytesIO(raw)).paragraphs[0].text == "Prueba del despliegue"
        elif format == "pdf":
            assert raw.startswith(b"%PDF")
        else:
            sheet = load_workbook(io.BytesIO(raw)).active
            assert sheet["B2"].value == "'=1+1"
        print(f"MCP {format}: archivo generado y validado en el servidor")
    for name in ("getEmployees", "getStockInfo"):
        assert not (await executor.execute(name, {}, ["customer"]))["audit"]["allowed"]
    print("MCP: controles de acceso correctos")


if __name__ == "__main__":
    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "chat-service"))
    verify_http()
    asyncio.run(verify_mcp())
