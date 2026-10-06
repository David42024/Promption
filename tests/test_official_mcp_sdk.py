"""Role-scoped business tools use the official MCP protocol implementation."""
import asyncio
import sys
from pathlib import Path

CHAT_SERVICE = Path(__file__).resolve().parents[1] / "chat-service"
if str(CHAT_SERVICE) not in sys.path:
    sys.path.insert(0, str(CHAT_SERVICE))

from app.mcp_tools import MCPToolExecutor


def test_sdk_clients_only_see_their_role_catalog():
    executor = MCPToolExecutor()

    async def names(role):
        return {tool.name for tool in await executor.available([role], True)}

    async def check():
        assert await names("customer") == {"getBrandInfo", "getShippingPolicy", "getCatalogSummary", "make_document"}
        brand = await executor.execute("getBrandInfo", {}, ["customer"])
        assert brand["result"]["brand"]["name"] == "Promption Shop"
        denied = await executor.execute("getEmployees", {}, ["customer"])
        assert denied["audit"]["allowed"] is False
        sales = await names("ventas")
        assert "getMarketingCampaigns" in sales
        assert "getEmployees" not in sales
        admin = await names("admin")
        assert "getEmployees" in admin
        assert "getInternalSecrets" not in admin

    asyncio.run(check())


def test_guest_can_read_only_explicit_public_mcp_data():
    executor = MCPToolExecutor()

    async def check():
        names = {tool.name for tool in await executor.available(["guest"], False)}
        assert names == {"getBrandInfo", "getShippingPolicy", "getCatalogSummary"}
        shipping = await executor.execute("getShippingPolicy", {}, ["guest"], authenticated=False)
        assert shipping["audit"]["allowed"]
        assert shipping["result"]["envios"]["canarias_ceuta_melilla"].startswith("5-7 días")
        for name in ("getEmployees", "getStockInfo", "make_document"):
            denied = await executor.execute(name, {}, ["guest"], authenticated=False)
            assert denied["audit"]["allowed"] is False

    asyncio.run(check())


def test_mcp_generates_office_and_pdf_files_without_formula_execution():
    import base64
    import io

    from docx import Document
    from openpyxl import load_workbook

    executor = MCPToolExecutor()

    async def check():
        for format in ("docx", "pdf", "xlsx"):
            response = await executor.execute("make_document", {
                "title": "Informe", "content": "Nombre,Valor\nDato,=1+1", "format": format,
            }, ["customer"])
            assert response["audit"]["allowed"]
            attachment = response["result"]
            assert attachment["name"] == f"Informe.{format}"
            data = base64.b64decode(attachment["data"])
            if format == "docx":
                assert Document(io.BytesIO(data)).paragraphs[0].text == "Informe"
            elif format == "pdf":
                assert data.startswith(b"%PDF")
            else:
                sheet = load_workbook(io.BytesIO(data)).active
                assert sheet["B2"].value == "'=1+1"
                assert sheet["B2"].data_type != "f"

    asyncio.run(check())
